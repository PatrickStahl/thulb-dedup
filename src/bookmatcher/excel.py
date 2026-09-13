from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Book
from .normalization import SEARCH_TITLE_COLUMN, VOLUME_COLUMN, clean_title_for_search

AUTHOR_COLUMN = "Autor/Herausgeber"
YEAR_COLUMN = "Erscheinungsjahr"
TITLE_COLUMN = "Titel (Auflage)"
ANNOTATION_COLUMN = "Anzahl des Exemplares in Thulb"
SOURCE_ROW_COLUMN = "Quellzeile"
PPN_COLUMN = "PPN"
CANDIDATE_PPN_COLUMN = "candidate_ppn"

REQUIRED_COLUMNS = {
    AUTHOR_COLUMN,
    YEAR_COLUMN,
    TITLE_COLUMN,
    ANNOTATION_COLUMN,
}


@dataclass(slots=True)
class SplitBooks:
    annotated: pd.DataFrame
    unannotated: pd.DataFrame


def load_excel(path: str | Path, sheet_name: str | int = 0) -> pd.DataFrame:
    """Load the source spreadsheet and retain the original source row number.

    XLSX/XLSM and CSV files are supported. Fully empty rows are removed.
    ``Quellzeile`` points to the row number in the original file (header is row
    1). If an input file already contains ``Quellzeile``, it is preserved.
    """
    frame = _read_table(path, sheet_name=sheet_name)

    return _prepare_source_frame(frame)


def _read_table(path: str | Path, sheet_name: str | int = 0) -> pd.DataFrame:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        return pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig")

    if suffix in {".xlsx", ".xlsm"}:
        return pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")

    raise ValueError(
        "Nicht unterstütztes Eingabeformat. Unterstützt werden: .csv, .xlsx, .xlsm"
    )


def _prepare_source_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()

    missing = REQUIRED_COLUMNS.difference(frame.columns)
    if missing:
        missing_text = ", ".join(sorted(missing))
        raise ValueError(f"Fehlende Pflichtspalten: {missing_text}")

    # pandas index 0 corresponds to source row 2 because row 1 is the header.
    if SOURCE_ROW_COLUMN not in frame.columns:
        frame.insert(0, SOURCE_ROW_COLUMN, frame.index + 2)

    source_columns = [column for column in frame.columns if column != SOURCE_ROW_COLUMN]
    frame = frame.dropna(how="all", subset=source_columns).copy()

    return frame


def reconstruct_original_output(
    input_path: str | Path,
    matches_path: str | Path,
    output_path: str | Path,
    *,
    sheet_name: str | int = 0,
    ppn_column: str = PPN_COLUMN,
) -> Path:
    """Write the original table layout enriched with match counts and PPNs.

    Rows keep the same columns and order as the original input, with one PPN
    column appended at the end. When one source row has multiple candidate PPNs,
    they are written as a comma-separated list.
    """
    input_frame = _read_table(input_path, sheet_name=sheet_name)
    keep_source_row = SOURCE_ROW_COLUMN in input_frame.columns
    source_frame = _prepare_source_frame(input_frame)
    matches = pd.read_csv(
        matches_path,
        sep=None,
        engine="python",
        encoding="utf-8-sig",
    )

    output = append_match_results_to_original(
        source_frame,
        matches,
        keep_source_row=keep_source_row,
        ppn_column=ppn_column,
    )

    output_path = Path(output_path)
    _write_table(output, output_path)

    return output_path


def cleanup_booklist(
    input_path: str | Path,
    output_path: str | Path,
    *,
    sheet_name: str | int = 0,
) -> Path:
    """Write a source table with inspectable search helper columns appended."""
    frame = load_excel(input_path, sheet_name=sheet_name)

    for column in (SEARCH_TITLE_COLUMN, VOLUME_COLUMN):
        if column in frame.columns:
            frame = frame.drop(columns=[column])

    search_titles: list[str] = []
    volumes: list[str] = []

    for title in frame[TITLE_COLUMN]:
        if pd.isna(title):
            search_titles.append("")
            volumes.append("")
            continue

        cleaned = clean_title_for_search(str(title))
        search_titles.append(cleaned.search_title)
        volumes.append(cleaned.volume or "")

    frame[SEARCH_TITLE_COLUMN] = search_titles
    frame[VOLUME_COLUMN] = volumes

    output_path = Path(output_path)
    _write_table(frame, output_path)

    return output_path


def append_match_results_to_original(
    source_frame: pd.DataFrame,
    matches: pd.DataFrame,
    *,
    keep_source_row: bool = False,
    ppn_column: str = PPN_COLUMN,
) -> pd.DataFrame:
    """Return the original rows with match count in the annotation column.

    ``matches`` is expected to be the output of ``bookmatcher.batch``. The
    result contains one row per original input row. Multiple candidate PPNs are
    joined in the appended PPN column.
    """
    missing = {SOURCE_ROW_COLUMN, CANDIDATE_PPN_COLUMN}.difference(matches.columns)
    if missing:
        missing_text = ", ".join(sorted(missing))
        raise ValueError(f"Fehlende Ergebnis-Spalten: {missing_text}")

    source_frame = source_frame.copy()
    if SOURCE_ROW_COLUMN not in source_frame.columns:
        source_frame.insert(0, SOURCE_ROW_COLUMN, source_frame.index + 2)

    matches_by_source_row = _group_candidate_ppns(matches)
    output_rows: list[dict[str, object]] = []

    for _, source_row in source_frame.iterrows():
        row = source_row.to_dict()
        source_key = _source_row_key(row.get(SOURCE_ROW_COLUMN))
        ppns = matches_by_source_row.get(source_key)

        if ppns is None:
            row[ppn_column] = ""
            output_rows.append(row)
            continue

        row[ANNOTATION_COLUMN] = len(ppns)

        row[ppn_column] = ", ".join(ppns)
        output_rows.append(row)

    frame_columns = list(source_frame.columns)
    if ppn_column not in frame_columns:
        frame_columns.append(ppn_column)

    output = pd.DataFrame(output_rows, columns=frame_columns)
    output_columns = _output_columns(
        source_frame.columns,
        keep_source_row=keep_source_row,
        ppn_column=ppn_column,
    )

    return output.loc[:, output_columns]


def _group_candidate_ppns(matches: pd.DataFrame) -> dict[int, list[str]]:
    grouped: dict[int, list[str]] = {}

    for _, match in matches.iterrows():
        source_key = _source_row_key(match.get(SOURCE_ROW_COLUMN))
        if source_key is None:
            continue

        grouped.setdefault(source_key, [])

        ppn = _optional_text(match.get(CANDIDATE_PPN_COLUMN))
        if ppn is not None:
            grouped[source_key].append(ppn)

    return grouped


def _source_row_key(value: object) -> int | None:
    if pd.isna(value):
        return None

    text = str(value).strip()
    if not text:
        return None

    try:
        return int(float(text))
    except ValueError:
        return None


def _output_columns(
    source_columns: Iterable[str],
    *,
    keep_source_row: bool,
    ppn_column: str,
) -> list[str]:
    columns = [
        column
        for column in source_columns
        if keep_source_row or column != SOURCE_ROW_COLUMN
    ]

    columns = [column for column in columns if column != ppn_column]
    columns.append(ppn_column)

    return columns


def _write_table(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        return

    if suffix == ".xlsx":
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            frame.to_excel(writer, index=False)
            worksheet = writer.sheets["Sheet1"]
            _style_output_sheet(worksheet)
        return

    raise ValueError(
        "Nicht unterstütztes Ausgabeformat. Unterstützt werden: .csv, .xlsx"
    )


def _style_output_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="000000")
    header_font = Font(color="FFFFFF", bold=True)
    alternate_fill = PatternFill("solid", fgColor="D9D9D9")
    border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    for column_index, column_cells in enumerate(
        worksheet.iter_cols(),
        start=1,
    ):
        if column_index % 2 == 1:
            for cell in column_cells[1:]:
                cell.fill = alternate_fill

        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 10),
            20,
        )

        for cell in column_cells:
            cell.border = border


def is_annotated(value: object) -> bool:
    """A row is annotated when the ThULB annotation cell is non-empty.

    This intentionally does not interpret the annotation. Both ``x`` and a
    value such as ``einmal verfügbar`` count as annotated. Positive/negative
    interpretation belongs to the later evaluation step.
    """
    if pd.isna(value):
        return False
    return bool(str(value).strip())


def split_by_annotation(
    frame: pd.DataFrame,
    annotation_column: str = ANNOTATION_COLUMN,
) -> SplitBooks:
    if annotation_column not in frame.columns:
        raise ValueError(f"Annotationsspalte nicht gefunden: {annotation_column}")

    mask = frame[annotation_column].map(is_annotated)
    return SplitBooks(
        annotated=frame.loc[mask].copy(),
        unannotated=frame.loc[~mask].copy(),
    )


def select_source_rows(
    frame: pd.DataFrame,
    source_rows: Iterable[int] | None,
) -> pd.DataFrame:
    """Optionally select rows by their original source row number.

    ``None`` means: keep all rows.
    """
    if source_rows is None:
        return frame.copy()

    selected = set(source_rows)
    return frame.loc[frame[SOURCE_ROW_COLUMN].isin(selected)].copy()


def export_split(
    split: SplitBooks,
    output_dir: str | Path,
    *,
    annotated_filename: str = "annotiert.csv",
    unannotated_filename: str = "unannotiert.csv",
) -> tuple[Path, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    annotated_path = output_dir / annotated_filename
    unannotated_path = output_dir / unannotated_filename

    split.annotated.to_csv(annotated_path, index=False)
    split.unannotated.to_csv(unannotated_path, index=False)

    return annotated_path, unannotated_path


def to_books(frame: pd.DataFrame) -> list[Book]:
    """Convert extracted rows to the minimal domain model for later matching."""
    books: list[Book] = []

    for _, row in frame.iterrows():
        books.append(
            Book(
                source_row=int(row[SOURCE_ROW_COLUMN]),
                author=_optional_text(row[AUTHOR_COLUMN]),
                title=_optional_text(row[TITLE_COLUMN]),
                year_raw=_optional_scalar(row[YEAR_COLUMN]),
                annotation_raw=_optional_text(row[ANNOTATION_COLUMN]),
            )
        )

    return books


def _optional_text(value: object) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _optional_scalar(value: object) -> int | float | str | None:
    if pd.isna(value):
        return None
    return value
